import copy, json, os
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph
from PIL import Image

ROOT = os.path.dirname(os.path.abspath(__file__))
EV = os.path.join(ROOT, "evidencia")
OUT = os.path.join(ROOT, "anexos_evidencias" + os.environ.get("DOCX_SUFFIX", "") + ".docx")
casos = json.load(open(os.path.join(ROOT, "casos.json"), encoding="utf8"))

# (archivo relativo a evidencia/, paso, nombre del paso)
CAP = {
    "CP-01": [("CP-01_paso1_home_medico.png", "Paso 1", "Inicio de sesión de MED1 en la interfaz: redirige a /home con el menú del médico"),
              ("CP-01_paso4_antes_de_cerrar.png", "Paso 4", "Sesión activa con datos visibles antes de cerrar sesión"),
              ("CP-01_paso4_atras_1.png", "Paso 4", "Botón Atrás #1 tras \"Cerrar sesión\": vuelve a /login sin datos"),
              ("CP-01_paso4_sin_sesion_pacientes.png", "Paso 4", "Sin sesión, abrir /pacientes directamente: redirige a /login (el tercer Atrás llega a about:blank, el inicio del historial de la pestaña, y no tiene captura)")],
    "CP-02": [("CP-02_paso1_ui_error.png", "Pasos 1 y 2", "Login con contraseña incorrecta: mensaje \"Credenciales inválidas\" y permanece en /login")],
    "CP-03": [("CP-03_paso1_registro_f.png", "Paso 1", "Registro de qa_registro01 con la contraseña del vector f"),
              ("CP-03_paso3_pacientes.png", "Paso 3", "Cuenta nueva (rol paciente) abre /pacientes"),
              ("CP-03_paso3_usuarios.png", "Paso 3", "Cuenta nueva abre /usuarios: \"No autorizado (solo administrador)\""),
              ("CP-03_paso3_auditoria.png", "Paso 3", "Cuenta nueva abre /auditoria: \"No autorizado (solo administrador)\""),
              ("CP-03_paso6_ui_vector_a.png", "Paso 6", "Registro desde el formulario con el vector a (11 caracteres): la cuenta se crea")],
    "CP-04": [("CP-04_paso9_ui_examen_ajeno.png", "Paso 9", "PAC_A abre /examenes/{EX_B}: la interfaz no muestra datos del examen del paciente B")],
    "CP-05": [("CP-05_paso1_intento1_error.png", "Paso 1", "Registro de Nuevo 1 con los datos de la Tabla 8: error de validación (Estado Civil)"),
              ("CP-05_paso1_intento2.png", "Paso 1", "Registro de Nuevo 1 eligiendo un Estado Civil: paciente registrado"),
              ("CP-05_paso5_editar.png", "Paso 5", "Edición de Camila Rojas: teléfono y dirección")],
    "CP-06": [("CP-06_paso1_listado.png", "Paso 1", "Listado de usuarios (1-10 de 14)"),
              ("CP-06_paso2_form.png", "Paso 2", "Edición de qa_rol01: correo y rol médico"),
              ("CP-06_paso3_menu_qa_rol01.png", "Paso 3", "Menú de qa_rol01 tras el cambio de rol (Pacientes y Exámenes Médicos)"),
              ("CP-06_paso5_dialogo.png", "Paso 5", "Diálogo \"Confirmar eliminación\" de qa_eliminar")],
    "CP-07": [("CP-07_paso1_form.png", "Paso 1", "Formulario Nuevo Examen con los datos de prueba"),
              ("CP-07_paso3_mis_resultados.png", "Paso 3", "\"Mis Resultados\" de PAC_A con los exámenes nuevos"),
              ("CP-07_paso4_editar.png", "Paso 4", "Edición de EX_A2 (diagnóstico y tratamiento)"),
              ("CP-07_paso4_tras_subir_v2.png", "Paso 4", "Formulario de EX_A2 después de subir resultado_examen_v2.pdf")],
    "CP-08": [("CP-08_paso1_detalle_med1.png", "Paso 1", "Detalle de EX_DEL con MED1: sin botón Eliminar"),
              ("CP-08_paso2_dialogo.png", "Paso 2", "ADM: diálogo \"Confirmar eliminación\" del examen")],
    "CP-09": [("CP-09_paso1_ui.png", "Paso 1", "Subida de resultado_examen.pdf a EX_A3 desde la interfaz")],
    "CP-10": [("CP-10_paso1_mis_resultados.png", "Paso 1", "\"Mis Resultados\" de PAC_A con la lista de exámenes"),
              ("CP-10_paso2_detalle.png", "Paso 2", "Detalle de EX_A1 (primera ejecución)"),
              ("CP-10_paso2_detalle_examen_2.png", "Paso 2", "Detalle del examen 2 (repetición, 7 de 7 detalles revisados)"),
              ("CP-10_paso2_detalle_examen_3.png", "Paso 2", "Detalle del examen 3 (repetición)")],
    "CP-11": [("CP-11_paso5_pantalla_auditoria.png", "Paso 5", "Pantalla Auditoría del administrador con los registros de las 13 acciones")],
    "CP-12": [("CP-12_paso3_pacientes_nuevo.png", "Paso 3", "PAC_A abre /pacientes/nuevo: se muestra el formulario"),
              ("CP-12_paso3_examenes.png", "Paso 3", "PAC_A abre /examenes"),
              ("CP-12_paso3_usuarios.png", "Paso 3", "PAC_A abre /usuarios: \"No autorizado\""),
              ("CP-12_paso3_rol_manipulado.png", "Paso 3", "user.rolNombre = admin en localStorage: la pantalla se abre, la API responde 403"),
              ("CP-12_ui_ADM_usuarios.png", "Pasos 1 y 2", "ADM: pantalla de usuarios"), ("CP-12_ui_ADM_pacientes.png", "Pasos 1 y 2", "ADM: pantalla de pacientes"),
              ("CP-12_ui_ADM_examenes.png", "Pasos 1 y 2", "ADM: pantalla de exámenes"), ("CP-12_ui_ADM_detalle.png", "Pasos 1 y 2", "ADM: detalle de examen con Eliminar y Subir Documento"),
              ("CP-12_ui_MED_usuarios.png", "Pasos 1 y 2", "MED1: /usuarios, \"No autorizado\""), ("CP-12_ui_MED_pacientes.png", "Pasos 1 y 2", "MED1: pantalla de pacientes"),
              ("CP-12_ui_MED_pacientes_con_eliminar.png", "Pasos 1 y 2", "MED1: la lista de pacientes muestra botones Eliminar (la API responde 403)"),
              ("CP-12_ui_MED_examenes.png", "Pasos 1 y 2", "MED1: pantalla de exámenes (sin Eliminar)"), ("CP-12_ui_MED_detalle.png", "Pasos 1 y 2", "MED1: detalle de examen (sin Eliminar)"),
              ("CP-12_ui_PAC_usuarios.png", "Pasos 1 y 2", "PAC_A: /usuarios, \"No autorizado\""), ("CP-12_ui_PAC_pacientes.png", "Pasos 1 y 2", "PAC_A: /pacientes con toast de error"),
              ("CP-12_ui_PAC_examenes.png", "Pasos 1 y 2", "PAC_A: pantalla de exámenes (solo \"Ver detalles\")"), ("CP-12_ui_PAC_detalle.png", "Pasos 1 y 2", "PAC_A: detalle de examen sin botones de gestión")],
    "CP-13": [("CP-13_paso1_registro.png", "Pasos 1 y 2", "Formulario /registro: sin casilla de aceptación ni enlace a la política"),
              ("CP-13_paso3_politica.png", "Paso 3", "Política de privacidad (se abre solo con sesión iniciada)")],
    "CP-14": [("img/CP-14_salida.png", "Pasos 1 a 5", "Extracto de la salida registrada (render del texto de CP-14.txt): contraseña vacía, 15 intentos, tokens inválidos y arranque sin JWT_SECRET")],
    "CP-15": [("img/CP-15_salida.png", "Pasos 1 a 6", "Extracto de la salida registrada (render de CP-15.txt): PUT ajenos rechazados y registro con RUT ajeno")],
    "CP-16": [("img/CP-16_salida.png", "Pasos 1 a 8", "Extracto de la salida registrada (render de CP-16.txt): escalamiento y accesos sin token")],
    "CP-17": [("CP-17_paso6_ui.png", "Paso 6", "Vector V1 desde la pantalla de login: \"Credenciales inválidas\"")],
    "CP-18": [("CP-18_N_chromium_detalle_examen.png", "Paso 4", "Detalle del examen con P1 y P3 en Chromium 156.0.8078.4 (Playwright 1.64.1, N)"),
              ("CP-18_N_firefox_detalle_examen.png", "Paso 4", "Detalle del examen en Firefox 157.0 (Playwright 1.64.1, N)"),
              ("CP-18_N_webkit_detalle_examen.png", "Paso 4", "Detalle del examen en WebKit 27.2 (Playwright 1.64.1, N)"),
              ("CP-18_N_msedge_detalle_examen.png", "Paso 4", "Detalle del examen en Edge real 156.0.4314.42"),
              ("CP-18_N-1_chromium_detalle_examen.png", "Paso 4", "Detalle del examen en Chromium 153.0.8010.12 (Playwright 1.63.0, N-1)"),
              ("CP-18_N-1_firefox_detalle_examen.png", "Paso 4", "Detalle del examen en Firefox 155.0 (Playwright 1.63.0, N-1)"),
              ("CP-18_N-1_webkit_detalle_examen.png", "Paso 4", "Detalle del examen en WebKit 26.6 (Playwright 1.63.0, N-1)"),
              ("CP-18_chromium_auditoria.png", "Paso 5", "Auditoría con el nombre del paciente como texto (Chromium 143, primera ejecución)"),
              ("CP-18_chromium_lista_pacientes.png", "Paso 3", "Lista de pacientes con el payload como texto (Chromium 143, primera ejecución)")],
    "CP-19": [("CP-19_paso5_ui_exe.png", "Paso 5", "Subida de script_malicioso.exe desde la interfaz: el diálogo no envía el archivo"),
              ("img/CP-19_salida.png", "Pasos 3 a 6", "Extracto de la salida de Newman (render de resumen_cp19.txt)")],
    "CP-20": [("img/CP-20_salida.png", "Pasos 1 a 5", "Extracto de la salida registrada (render de CP-20.txt)")],
    "CP-21": [("img/CP-21_salida.png", "Paso 1", "Extracto de testssl.sh contra el proxy TLS del equipo (protocolos, HSTS, cadena de confianza)")],
    "CP-22": [("img/CP-22_salida.png", "Pasos 2 a 7", "Extracto de la salida registrada (render de CP-22.txt)")],
    "CP-25": [("cp25/img/base_1usuario_cli.png", "Medición base", "k6 con 1 usuario durante 1 minuto (salida de k6, render HTML del texto real)"),
              ("cp25/img/A_run1_cli.png", "Parte A", "Carga normal, repetición 1: 20 usuarios por 3 minutos (salida de k6)"),
              ("cp25/img/A_run2_cli.png", "Parte A", "Carga normal, repetición 2 (salida de k6)"),
              ("cp25/img/A_run3_cli.png", "Parte A", "Carga normal, repetición 3 (salida de k6)"),
              ("cp25/img/B_run1_cli.png", "Parte B", "Carga máxima, repetición 1: rampa a 200 usuarios y 5 minutos sostenidos (salida de k6)"),
              ("cp25/img/B_run2_cli.png", "Parte B", "Carga máxima, repetición 2 (salida de k6)"),
              ("cp25/img/B_run3_cli.png", "Parte B", "Carga máxima, repetición 3 (salida de k6)"),
              ("cp25/img/C_volumen_cli.png", "Parte C", "Volumen: 50 consultas con el token del médico y 50 con el del paciente, 200 exámenes cada una (salida de k6)"),
              ("cp25/img/CP-25_parteC_mis_resultados.png", "Parte C", "\"Mis Resultados\" del paciente con 200 exámenes (captura real)"),
              ("cp25/img/lighthouse_cli.png", "Parte D", "Lighthouse 13.5.0, móvil con 4G y caché vacía, sobre /login y /home (salida del script)"),
              ("cp25/img/E_documentos_cli.png", "Parte E", "Documentos: 10 subidas y 10 descargas de 2 MB (salida de k6)"),
              ("cp25/img/F_estres_cli.png", "Parte F", "Estrés: de 20 a 400 usuarios y vuelta a 20 (salida de k6)"),
              ("cp25/img/CP-25_estres_detalle.png", "Parte F", "Gráfico: usuarios virtuales, solicitudes por segundo, CPU y latencia p95 durante el estrés"),
              ("cp25/img/G_sin_token_cli.png", "Parte G", "Respuesta 401 sin token: 20 solicitudes (salida de k6)"),
              ("cp25/img/CP-25_hardware_global.png", "Todas las partes", "Gráfico: CPU y memoria de la API y de PostgreSQL a lo largo de todas las cargas (docker stats cada 8 s)")],
    "CP-23": [("img/CP-23_salida.png", "Pasos 1 a 6", "Extracto de la salida registrada (render de CP-23.txt)")],
}

d = docx.Document(os.path.join(ROOT, "..", "docs", "EP3_ISY1102_Estudiante_Formato_Informe.docx"))
body = d.element.body
kids = list(body.iterchildren())
h1 = next(copy.deepcopy(k) for k in kids if k.tag == qn("w:p") and "".join(x.text or "" for x in k.iter(qn("w:t"))).startswith("5. Ejecución de pruebas"))
h2 = next(copy.deepcopy(k) for k in kids if k.tag == qn("w:p") and "".join(x.text or "" for x in k.iter(qn("w:t"))).startswith("5.2 "))
for k in kids:
    if k.tag != qn("w:sectPr"):
        body.remove(k)
sect = body.find(qn("w:sectPr"))


def set_text(p, text):
    runs = p.findall(qn("w:r"))
    proto = copy.deepcopy(runs[0])
    for r in runs:
        p.remove(r)
    for t in proto.findall(qn("w:t")):
        proto.remove(t)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    proto.append(t)
    p.append(proto)


def heading(proto, text):
    h = copy.deepcopy(proto)
    set_text(h, text)
    sect.addprevious(h)


def para(text, bold=False, size=None, italic=False):
    p = OxmlElement("w:p")
    sect.addprevious(p)
    par = Paragraph(p, d._body)
    r = par.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    par.paragraph_format.space_after = Pt(4)
    par.paragraph_format.keep_with_next = bold
    return par


def picture(path):
    w, h = Image.open(path).size
    width = Cm(16.4)
    height = width * h / w
    if height > Cm(21):
        width = Cm(21) * w / h
    p = OxmlElement("w:p")
    sect.addprevious(p)
    pp = Paragraph(p, d._body)
    pp.add_run().add_picture(path, width=int(width))
    pp.paragraph_format.keep_with_next = True


heading(h1, "Anexo de evidencias por caso de prueba")
para("Imágenes de la ejecución de los casos CP-01 a CP-23 (fases 1 y 2) y de CP-25 (fase 3). Cada capítulo corresponde a un caso del informe y cada figura indica el paso del caso y qué muestra. Las imágenes con la leyenda \"extracto de la salida registrada\" son un render HTML del texto real de la salida guardada en resultados/evidencia/, porque esos pasos se ejecutaron por API o SQL y no tienen pantalla. El resultado de cada caso está en la sección 5.2 del documento seccion5.docx.", size=10, italic=True)
total = 0
for cid in sorted(CAP):
    c = casos[cid]
    heading(h2, "%s · %s" % (cid, c["nombre"]))
    para("Resultado: %s" % c["resultado"].split("<br>")[0].split(" (")[0], bold=True, size=10)
    n = 0
    for rel, paso, nombre in CAP[cid]:
        path = os.path.join(EV, rel)
        if not os.path.exists(path):
            print("FALTA", rel)
            continue
        n += 1
        total += 1
        para("%s. %s" % (paso, nombre), bold=True, size=10)
        picture(path)
        para("Archivo: resultados/evidencia/" + rel, size=8, italic=True)
    print(cid, n, "imágenes")
d.save(OUT)
json.dump({k: len(v) for k, v in CAP.items()}, open(os.path.join(ROOT, "imagenes_por_caso.json"), "w"))
print("guardado", OUT, "| imágenes", total)
