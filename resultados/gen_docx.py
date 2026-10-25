import copy, json, os, re, sys
import docx
from docx.shared import Cm
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(ROOT, "..", "docs", "EP3_ISY1102_Estudiante_Formato_Informe.docx")
OUT = os.path.join(ROOT, "seccion5" + os.environ.get("DOCX_SUFFIX", "") + ".docx")
casos = json.load(open(os.path.join(ROOT, "casos.json"), encoding="utf8"))
extra = json.load(open(os.path.join(ROOT, "anexos.json"), encoding="utf8"))

d = docx.Document(TEMPLATE)
body = d.element.body
children = list(body.iterchildren())


def txt(el):
    return "".join(x.text or "" for x in el.iter(qn("w:t")))


# ---- localizar la sección 5 en la plantilla ----
i5 = next(i for i, el in enumerate(children) if el.tag == qn("w:p") and txt(el).startswith("5. Ejecución de pruebas"))
for el in children[:i5]:
    body.remove(el)
children = list(body.iterchildren())
def find(prefix):
    return next(el for el in children if el.tag == qn("w:p") and txt(el).startswith(prefix))

h2_proto = next(el for el in children if el.tag == qn("w:p") and txt(el).startswith("5.2 "))
tbls = [el for el in children if el.tag == qn("w:tbl")]
t52, t53scale, t53inc = tbls


def set_text_runs(p, text, size=None, bold=None):
    """Deja en el párrafo un único run con el formato del primer run existente."""
    runs = p.findall(qn("w:r"))
    proto = copy.deepcopy(runs[0]) if runs else None
    for r in runs:
        p.remove(r)
    if proto is None:
        proto = OxmlElement("w:r")
    for t in proto.findall(qn("w:t")):
        proto.remove(t)
    for br in proto.findall(qn("w:br")):
        proto.remove(br)
    rpr = proto.find(qn("w:rPr"))
    if rpr is not None:
        if size is not None:
            for tag in ("w:sz", "w:szCs"):
                e = rpr.find(qn(tag))
                if e is None:
                    e = OxmlElement(tag)
                    rtl = rpr.find(qn("w:rtl"))
                    if rtl is not None:
                        rtl.addprevious(e)
                    else:
                        rpr.append(e)
                e.set(qn("w:val"), str(size))
        if bold is not None:
            for tag in ("w:b", "w:bCs"):
                e = rpr.find(qn(tag))
                if e is not None:
                    e.set(qn("w:val"), "1" if bold else "0")
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    proto.append(t)
    p.append(proto)


def set_cell(tc, text, size=None, bold=None):
    ps = tc.findall(qn("w:p"))
    first = ps[0]
    for p in ps[1:]:
        tc.remove(p)
    lines = text.split("<br>") if text else [""]
    set_text_runs(first, lines[0], size, bold)
    prev = first
    for ln in lines[1:]:
        np_ = copy.deepcopy(first)
        set_text_runs(np_, ln, size, bold)
        prev.addnext(np_)
        prev = np_


def set_grid(tbl, widths):
    grid = tbl.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    tw = tbl.find(qn("w:tblPr")).find(qn("w:tblW"))
    tw.set(qn("w:w"), str(sum(widths)))
    tw.set(qn("w:type"), "dxa")
    lay = tbl.find(qn("w:tblPr")).find(qn("w:tblLayout"))
    if lay is None:
        lay = OxmlElement("w:tblLayout")
        tbl.find(qn("w:tblPr")).append(lay)
    lay.set(qn("w:type"), "fixed")


def row_cells(tr):
    return tr.findall(qn("w:tc"))


EV = os.path.join(ROOT, "evidencia")
PICS = json.load(open(os.path.join(ROOT, "imagenes.json"), encoding="utf8"))


IMG_N = json.load(open(os.path.join(ROOT, "imagenes_por_caso.json")))


def add_pics(tc, cid):
    """Las imágenes de los casos viven en anexos_evidencias.docx; aquí solo se deja la referencia."""
    n = IMG_N.get(cid)
    if not n:
        return
    p = copy.deepcopy(tc.findall(qn("w:p"))[0])
    set_text_runs(p, "Imágenes: anexos_evidencias.docx, capítulo %s (%d %s)." % (cid, n, "imagen" if n == 1 else "imágenes"), size=16)
    tc.append(p)


FILL = {"CRITICO": "ff0000", "GRAVE": "ffc000", "MODERADO": "ffff00", "LEVE": None}


def paint(tc, fill):
    pr = tc.find(qn("w:tcPr"))
    for e in pr.findall(qn("w:shd")):
        pr.remove(e)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        mar = pr.find(qn("w:tcMar"))
        if mar is not None:
            mar.addprevious(shd)
        else:
            pr.append(shd)


def build_table(proto_tbl, headers, rows, widths, group_key=None, size=16, paint_last=False, pics=None):
    """rows: lista de tuplas o ('__grupo__', texto). Usa filas de la plantilla (encabezado, grupo, caso)."""
    tbl = copy.deepcopy(proto_tbl)
    trs = tbl.findall(qn("w:tr"))
    head, grp, case = trs[0], trs[1], trs[2]
    for tr in trs:
        tbl.remove(tr)
    h = copy.deepcopy(head)
    for tc, t in zip(row_cells(h), headers):
        set_cell(tc, t)
    tbl.append(h)
    for r in rows:
        if r[0] == "__grupo__":
            g = copy.deepcopy(grp)
            set_cell(row_cells(g)[0], r[1])
            tbl.append(g)
        else:
            c = copy.deepcopy(case)
            for k, (tc, t) in enumerate(zip(row_cells(c), r)):
                set_cell(tc, t, size=size, bold=(True if k == 0 else None))
            if paint_last:
                paint(row_cells(c)[-1], FILL.get(r[-1]))
            if pics is not None and r[0] in IMG_N:
                add_pics(row_cells(c)[3], r[0])
            tbl.append(c)
    set_grid(tbl, widths)
    return tbl


# ---- 5.1 ----
ids = sorted(casos)
ejec = [c for c in ids if casos[c]["resultado"] != "sin ejecutar"]
aprob = [c for c in ejec if casos[c]["resultado"].startswith("aprobado")]
fall = [c for c in ejec if casos[c]["resultado"].startswith("falla")]
bloq = [c for c in ejec if casos[c]["resultado"].startswith("bloqueado")]
sinej = [c for c in ids if casos[c]["resultado"] == "sin ejecutar"]
assert len(aprob) + len(fall) + len(bloq) == len(ejec)
p_instr = find("Presentar una visión general del")
p_a, p_b, p_c = find("Cantidad total de casos de prue"), find("Cantidad de casos ejecutados"), find("Resultados obtenidos, cuantos ")
body.remove(p_instr)
set_text_runs(p_a, f"Cantidad total de casos de prueba diseñados: {len(ids)}")
def rangos(lst):
    n = [int(c[3:]) for c in lst]; out = []; s = p = n[0]
    for x in n[1:] + [None]:
        if x is not None and x == p + 1:
            p = x; continue
        out.append("CP-%02d" % s if s == p else "CP-%02d a CP-%02d" % (s, p)); s = p = x
    return " y ".join(out)
set_text_runs(p_b, f"Cantidad de casos ejecutados: {len(ejec)} ({rangos(ejec)})")
set_text_runs(p_c, f"Resultados obtenidos: pasaron {len(aprob)}, fallaron {len(fall)}, quedaron bloqueados {len(bloq)} y quedaron sin ejecutar {len(sinej)} (" + ", ".join(sinej) + ")")

# ---- 5.2 ----
for key in ["Detallar la ejecución caso a ca", "ID del caso de prueba", "Resultado de ejecución", "Evidencia", "Incidente reportado (en el ", "Fecha y persona que ejecuta l"]:
    body.remove(find(key))
rows52 = []
grupos = {}
for c in ids:
    g = casos[c]["grupo"]
    if g not in grupos:
        grupos[g] = True
        rows52.append(("__grupo__", f"TIPO DE PRUEBA: {g}"))
    rows52.append((c, casos[c]["nombre"], casos[c]["resultado"], casos[c]["evidencia"]))
new52 = build_table(t52, ["CASO", "DESCRIPCIÓN", "RESULTADO", "EVIDENCIA"], rows52, [800, 1600, 3300, 4100], pics=PICS)
t52.addprevious(new52)
body.remove(t52)

# ---- 5.3 ----
def label(c):
    first = casos[c]["resultado"].split("<br>")[0]
    if "brecha declarada" in first:
        return "falla (brecha declarada y falla esperada por diseño)"
    e_ = "esperada por diseño" in first
    n_ = "hallazgo nuevo" in first
    return "falla (" + ("falla esperada por diseño y hallazgo nuevo" if e_ and n_ else "falla esperada por diseño" if e_ else "hallazgo nuevo") + ")"


p_scale = next(el for el in body.iterchildren() if el.tag == qn("w:p") and txt(el).startswith("Establecer escala"))
set_text_runs(p_scale, "Escala de medición de la criticidad de los incidentes:")
rows53 = []
grupos = {}
for c in ids:
    if casos[c]["impacto"]:
        g = casos[c]["grupo"]
        if g not in grupos:
            grupos[g] = True
            rows53.append(("__grupo__", f"TIPO DE PRUEBA: {g}"))
        rows53.append((c, casos[c]["nombre"], label(c), casos[c]["impacto"]))
new53 = build_table(t53inc, ["CASO", "DESCRIPCIÓN", "RESULTADO", "IMPACTO"], rows53, [800, 3700, 3800, 1500], size=18, paint_last=True)
t53inc.addprevious(new53)
body.remove(t53inc)

set_grid(t53scale, [1500, 3800, 2600, 1900])

# ---- 5.4 y 5.5 ----
last_p = list(body.iterchildren())[-2]  # párrafo vacío tras la última tabla (antes de sectPr)


def add_h2(text, before):
    h = copy.deepcopy(h2_proto)
    set_text_runs(h, text)
    before.addprevious(h)
    return h


sect = body.find(qn("w:sectPr"))
add_h2("5.4 Casos que no se pudieron ejecutar y detalles", sect)
note = copy.deepcopy(p_scale)
set_text_runs(note, extra["nota_54"])
sect.addprevious(note)
sect.addprevious(build_table(t52, ["CASO", "DESCRIPCIÓN", "SITUACIÓN", "DETALLE"], [(a, b, c, e) for a, b, c, e in extra["no_ejecutados"]], [800, 1600, 2000, 5400]))
sp = copy.deepcopy(p_scale)
set_text_runs(sp, "Casos ejecutados con limitaciones o desviaciones respecto del informe:")
sect.addprevious(sp)
sect.addprevious(build_table(t52, ["CASO", "DESCRIPCIÓN", "TIPO", "DETALLE"], [tuple(r) for r in extra["limitaciones"]], [800, 1600, 1500, 5900]))
add_h2("5.5 Herramientas utilizadas y versiones", sect)
n2 = copy.deepcopy(p_scale)
set_text_runs(n2, extra["nota_55"])
sect.addprevious(n2)
sect.addprevious(build_table(t52, ["HERRAMIENTA (TABLA 6)", "VERSIÓN UTILIZADA", "USO EN LA EJECUCIÓN", "DIFERENCIA CON EL INFORME"], [("__grupo__", "HERRAMIENTAS DE LA TABLA 6")] + [tuple(r) for r in extra["tabla6"]] + [("__grupo__", "HERRAMIENTAS USADAS QUE NO FIGURAN EN LA TABLA 6")] + [tuple(r) for r in extra["adicionales"]], [1700, 2100, 2800, 3200]))
n3 = copy.deepcopy(p_scale)
set_text_runs(n3, extra["nota_sut"])
sect.addprevious(n3)

# ---- Anexos con imágenes de las herramientas ----
add_h2("Anexos. Resultados de las herramientas y evidencia de su uso", sect)
for title, desc, paths in json.load(open(os.path.join(ROOT, "anexos_visual.json"), encoding="utf8")):
    h = copy.deepcopy(p_scale)
    set_text_runs(h, title, bold=True)
    sect.addprevious(h)
    dp = copy.deepcopy(p_scale)
    set_text_runs(dp, desc)
    sect.addprevious(dp)
    for rel in paths:
        path = os.path.join(EV, rel)
        if not os.path.exists(path):
            print("FALTA imagen", rel)
            continue
        pp = OxmlElement("w:p")
        sect.addprevious(pp)
        Paragraph(pp, d._body).add_run().add_picture(path, width=Cm(16.4))
        cp = copy.deepcopy(p_scale)
        set_text_runs(cp, "Figura: resultados/evidencia/" + rel, size=16)
        sect.addprevious(cp)

d.save(OUT)
print("guardado", OUT, "| ejecutados", len(ejec), "aprobados", len(aprob), "fallan", len(fall), "sin ejecutar", len(sinej))
